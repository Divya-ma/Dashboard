"""The journal in the app: entry and exit forms, the Settings list of ideas, and the Archive's new columns and editors."""

from datetime import datetime, timedelta, timezone

import pytest

from core.archive import build_archive_rows
from core.journal_service import JournalService
from core.models import Contract, Leg, Structure, StructureStatus, StructureType, Trade, TradeEventType
from db.repository import Repository
from ui.callbacks import archive_callbacks as ac
from ui.callbacks import dissolve_callbacks as dcb
from ui.callbacks import settings_callbacks as sec
from ui.callbacks import structure_detail_callbacks as dc
from ui.callbacks import structures_callbacks as sc
from ui.container import Container
from ui.layouts import archive_tab as at
from ui.layouts.structure_detail import structure_detail_layout

GRID = ("active", "all", "name", {}, {})
LIVE_STORE = {"CLX26": {"price": 75.0}, "CLZ26": {"price": 74.0}}
T0 = datetime(2026, 10, 1, 9, 0, tzinfo=timezone.utc)


class FakeSnapshotter:
    def __init__(self):
        self.calls = []

    def snapshot(self, structure, when):
        self.calls.append((structure.name, when))
        n = len(self.calls)
        return {"at": when.isoformat(), "product": "CL", "curve_family": "outright", "prices": {"CO": 75.0 + n, "CL": 70.0 + n, "GO": 700.0 + n},
                "move_day": "Bullish", "move_min": "Stalemate", "curve": {"CL": "Contango", "BRN": "Backwardation", "G": "Contango - Backwardation"}, "atr": 0.5 * n}


def contract(symbol, month):
    return Contract(product="CL", contract_month=month, contract_year=2026, symbol=symbol, multiplier=1000.0, tick_size=0.01, tick_value=10.0)


def spread(name="spr", status=StructureStatus.SHELL):
    legs = [Leg(contract=contract("CLX26", 11), ratio=1), Leg(contract=contract("CLZ26", 12), ratio=-1)]
    return Structure(name=name, structure_type=StructureType.SPREAD, products=["CL"], legs=legs, status=status)


@pytest.fixture
def repo(tmp_db_path):
    return Repository(str(tmp_db_path))


@pytest.fixture
def snapshotter():
    return FakeSnapshotter()


@pytest.fixture(autouse=True)
def env(repo, snapshotter, monkeypatch):
    journal = JournalService(repo, snapshotter, clock=lambda: T0, background=False)
    fresh = Container(repository=repo, journal=journal)
    for module in (dc, sc, dcb, ac, sec):
        monkeypatch.setattr(module, "container", fresh)
    import ui.layouts.structure_detail as layout_module

    monkeypatch.setattr(layout_module, "container", fresh)
    return fresh


def save(repo, structure, trades=()):
    for leg in structure.legs:
        repo.save_contract(leg.contract)
    repo.save_structure(structure)
    for t in trades:
        repo.save_trade(t.model_copy(update={"structure_id": structure.structure_id, "leg_id": structure.legs[0].leg_id}))
    return structure.structure_id


# ---------- Settings ----------


def test_the_ideas_list_is_saved_from_settings(env):
    assert "2 trade idea" in sec.save_trade_ideas(1, "Mean reversion\nBreakout")
    assert env.journal.idea_options() == ["Mean reversion", "Breakout"]
    assert "0 trade idea" in sec.save_trade_ideas(1, "  ")
    with pytest.raises(Exception):
        sec.save_trade_ideas(None, "x")


# ---------- entering and exiting ----------


def test_an_entry_records_ideas_payups_and_the_market_snapshot(repo, env, snapshotter):
    env.journal.save_idea_options("Mean reversion\nBreakout")
    sid = save(repo, spread())
    out = dc.confirm_trade_entry(1, 1.25, 5, "buy", None, 1.0, 2.5, [], sid, LIVE_STORE, *GRID, ["Breakout", "Unknown"], 2, [1, None])
    assert out[-1] != "danger"
    journal = env.journal.journal(sid)
    assert journal["ideas"] == ["Breakout"] and journal["payups"]["entry"] == {"total": 2.0, "legs": {"CLX26": 1.0}}
    assert journal["entry"]["prices"]["CL"] == 71.0 and journal["exit"] is None
    # adding to the position: more payups and ideas, but the snapshot is still the first entry's
    dc.confirm_trade_entry(1, 1.30, 2, "buy", None, None, None, [], sid, LIVE_STORE, *GRID, ["Mean reversion"], 1, None)
    journal = env.journal.journal(sid)
    assert journal["ideas"] == ["Breakout", "Mean reversion"] and journal["payups"]["entry"]["total"] == 3.0 and journal["entry"]["prices"]["CL"] == 71.0
    # every entry is snapshotted for its own Archive row (first in, first out), but `entry` stays the FIRST entry's
    assert len(snapshotter.calls) == 2 and len(journal["entries"]) == 2


def test_a_failed_entry_leaves_the_journal_alone(repo, env, snapshotter):
    sid = save(repo, spread())
    out = dc.confirm_trade_entry(1, None, 5, "buy", None, None, None, [], sid, LIVE_STORE, *GRID, ["x"], 2, None)
    assert out[3] == "danger" or "danger" in out
    assert repo.get_journal(sid) is None and snapshotter.calls == []


def test_exits_record_payups_and_the_last_exit_snapshot(repo, env, snapshotter):
    sid = save(repo, spread())
    dc.confirm_trade_entry(1, 1.25, 10, "buy", None, None, None, [], sid, LIVE_STORE, *GRID)
    dc.execute_full_exit(1, 1.5, 4, None, sid, *GRID, 1, [None, 2])  # a partial exit
    journal = env.journal.journal(sid)
    assert journal["payups"]["exit"] == {"total": 1.0, "legs": {"CLZ26": 2.0}} and journal["exit"]["prices"]["CL"] == 72.0
    dc.execute_full_exit(1, 1.6, 6, None, sid, *GRID, 2, None)  # the rest
    journal = env.journal.journal(sid)
    assert journal["payups"]["exit"]["total"] == 3.0 and journal["exit"]["prices"]["CL"] == 73.0  # the last exit is the one kept
    assert repo.get_structure(sid).status == StructureStatus.CLOSED


def test_a_dissolve_takes_the_exit_snapshot(repo, env, snapshotter):
    a, b = spread("a"), spread("b")
    b = b.model_copy(update={"legs": [leg.model_copy(update={"ratio": -leg.ratio}) for leg in b.legs]})
    ids = [save(repo, s) for s in (a, b)]
    for sid in ids:
        dc.confirm_trade_entry(1, 1.25, 2, "buy", None, None, None, [], sid, LIVE_STORE, *GRID)
    snapshotter.calls.clear()
    form = (ids, [], [], [75.0, 74.0], [{"type": "dissolve-price", "index": "CLX26"}, {"type": "dissolve-price", "index": "CLZ26"}], "legs", None, None)
    dcb.execute(1, *form, *GRID)
    assert all(env.journal.journal(sid)["exit"] is not None for sid in ids) and len(snapshotter.calls) == 2


def test_the_forms_show_the_ideas_and_the_payup_boxes(repo, env):
    env.journal.save_idea_options("Carry")
    s = spread()
    save(repo, s)
    text = str(structure_detail_layout(s, {}, [], None))
    for expected in ("trade-entry-ideas", "Carry", "trade-entry-payups", "entry-leg-payup", "btn-toggle-entry-leg-payups"):
        assert expected in text, expected
    opened = s.model_copy(update={"status": StructureStatus.OPEN, "legs": [leg.model_copy(update={"lots": 3.0, "entry_price": 70.0}) for leg in s.legs]})
    text = str(structure_detail_layout(opened, {}, [], None))
    for expected in ("trade-exit-payups", "exit-leg-payup"):
        assert expected in text, expected


# ---------- the Archive ----------


def closed_structure(repo, name="done", day=0):
    s = spread(name, StructureStatus.CLOSED)
    s = s.model_copy(update={"legs": [leg.model_copy(update={"direction": "buy", "entry_price": 70.0, "average_entry_price": 70.0}) for leg in s.legs], "closed_at": T0 + timedelta(days=3)})
    trades = [Trade(structure_id="s", leg_id="l", event_type=TradeEventType.TRADE, lots=2, price=1.0, direction="buy", timestamp=T0 + timedelta(days=day),
                    stop_loss_price=0.5, target_price=2.0),
              Trade(structure_id="s", leg_id="l", event_type=TradeEventType.FULL_EXIT, lots=2, price=1.5, direction="sell", timestamp=T0 + timedelta(days=3), realized_pnl=1000.0)]
    return save(repo, s, trades)


def test_archive_rows_carry_the_journal_columns(repo, env):
    sid = closed_structure(repo)
    env.journal.save_idea_options("Carry\nBreakout")
    env.journal.set_ideas(sid, ["Carry", "Breakout"])
    env.journal.add_entry_inputs(sid, [], 2, {"CLX26": 1})
    env.journal.add_exit_payups(sid, 1, None)
    env.journal.record(repo.get_structure(sid), "entry", T0)
    env.journal.record(repo.get_structure(sid), "exit", T0 + timedelta(days=3), overwrite=True)
    (row,) = ac._load_rows()
    assert row["rr_text"] == "1 : 2.00" and row["ideas"] == ["Carry", "Breakout"] and row["payup_entry"] == 2.0 and row["payup_exit"] == 1.0
    assert row["e_px_cl"] == 71.0 and row["x_px_go"] == 702.0 and row["e_move_day"] == "Bullish" and row["x_move_min"] == "Stalemate"
    assert row["e_curve_brn"] == "Backwardation" and row["e_curve_go"] == "Contango - Backwardation" and row["x_atr"] == 1.0
    cells = at.build_table_rows([row])[0]
    assert cells["rr"] == "1 : 2.00" and cells["ideas"] == "Carry, Breakout" and cells["payup_entry"] == "2" and cells["e_px_co"] == 76.0
    columns = {c["id"] for c in at.TABLE_COLUMNS}
    assert {"rr", "ideas", "payup_entry", "payup_exit", "e_px_co", "e_curve_go", "x_atr", "x_move_day"} <= columns
    assert all(column in cells for column in columns - {"id"} if column != "id")
    detail = str(at.build_detail(row))
    assert "Risk : Reward 1 : 2.00" in detail and "At entry" in detail and "Contango - Backwardation" in detail and "CLX26 1" in detail


def test_a_trade_without_a_journal_shows_dashes(repo, env):
    closed_structure(repo, "plain")
    (row,) = ac._load_rows()
    cells = at.build_table_rows([row])[0]
    assert cells["e_px_cl"] == "—" and cells["ideas"] == "—" and cells["payup_entry"] == "—" and cells["rr"] == "1 : 2.00"
    assert "No market snapshot yet" in str(at.build_detail(row))


def test_the_archive_editor_saves_ideas_and_payups_and_keeps_the_leg_figures(repo, env):
    sid = closed_structure(repo)
    env.journal.save_idea_options("Carry\nBreakout")
    env.journal.add_entry_inputs(sid, ["Carry"], 1, {"CLX26": 1})
    options, ideas, entry, exit_, conviction, reason, mistakes, reason_options, mistake_options = ac.fill_journal_editor(sid)
    assert [o["value"] for o in options] == ["Carry", "Breakout"] and ideas == ["Carry"] and entry == 1.0 and exit_ is None
    assert ac.save_journal_edits(1, sid, ["Breakout", "Carry"], 4, 2) .startswith("Saved")
    journal = env.journal.journal(sid)
    assert journal["ideas"] == ["Breakout", "Carry"] and journal["payups"]["entry"] == {"total": 4.0, "legs": {"CLX26": 1.0}} and journal["payups"]["exit"]["total"] == 2.0
    ac.save_journal_edits(1, sid, [], "", None)
    assert env.journal.journal(sid)["ideas"] == [] and env.journal.journal(sid)["payups"]["entry"]["total"] is None
    assert ac.fill_journal_editor(None) == ([], [], None, None, None, None, [], [], []) and ac.save_journal_edits(1, None, [], None, None) == "Select a trade first."


def test_the_backfill_button_fills_old_trades_a_few_at_a_time(repo, env, snapshotter):
    for i in range(6):
        closed_structure(repo, f"old{i}")
    first = ac.backfill_market_data(1)
    assert f"Filled {ac.BACKFILL_PER_CLICK}" in first and "3 left" in first
    second = ac.backfill_market_data(1)
    assert "Filled 3" in second and "0 left" in second
    assert "already has" in ac.backfill_market_data(1)
    assert len(snapshotter.calls) == 6 * 2  # an entry and an exit snapshot each
    with pytest.raises(Exception):
        ac.backfill_market_data(None)


# ---------- Excel export ----------


def test_the_archive_exports_to_an_excel_workbook(repo, env):
    import io
    import json

    import pandas as pd
    from dash.exceptions import PreventUpdate

    a = closed_structure(repo, "first")
    closed_structure(repo, "second", day=1)
    env.journal.save_idea_options("Carry\nBreakout")
    env.journal.add_entry_inputs(a, ["Carry", "Breakout"], 2, {"CLX26": 1})
    env.journal.record(repo.get_structure(a), "entry", T0)
    rows = json.loads(json.dumps(ac._load_rows()))  # as the Archive tab keeps them in its store
    out = ac.export_to_excel(1, rows)
    assert out["filename"] == f"archive_{datetime.now().date().isoformat()}.xlsx"

    import base64

    content = out["content"]
    book = pd.read_excel(io.BytesIO(base64.b64decode(content)), sheet_name=None)
    assert list(book) == ["Trades", "Legs", "Summary"]
    trades = book["Trades"]
    assert len(trades) == 2 and {c["name"] for c in at.TABLE_COLUMNS} <= set(trades.columns)
    first = trades[trades["Structure Name"] == "first"].iloc[0]
    assert first["Risk : Reward"] == pytest.approx(2.0) and first["Trade Ideas"] == "Carry, Breakout" and first["Payups in Entry"] == 2
    assert first["Entry CL price"] == pytest.approx(71.0) and first["Entry CL curve"] == "Contango" and pd.isna(first["Exit CL price"])
    assert first["Realized PnL ($, net of TC)"] == pytest.approx(1000.0) and str(first["Entry Date"])[:10] == "2026-10-01"
    legs = book["Legs"]
    assert len(legs) == 4 and set(legs["Contract"]) == {"CLX26", "CLZ26"}
    summary = dict(zip(book["Summary"]["Item"], book["Summary"]["Value"]))
    assert summary["Trades"] == 2 and summary["Total PnL ($, net of TC)"] == pytest.approx(2000.0)
    # only the rows the table shows are exported
    one = ac.export_to_excel(1, rows[:1])
    assert len(pd.read_excel(io.BytesIO(base64.b64decode(one["content"])), sheet_name="Trades")) == 1
    with pytest.raises(PreventUpdate):
        ac.export_to_excel(1, [])
    with pytest.raises(PreventUpdate):
        ac.export_to_excel(None, rows)
    assert "archive-export-btn" in str(at.archive_layout()) and "archive-export-download" in str(at.archive_layout())


def test_the_workbook_has_money_formats_and_a_frozen_header(repo, env):
    import io

    from openpyxl import load_workbook

    from core.archive_export import archive_to_excel

    closed_structure(repo, "fmt")
    sheet = load_workbook(io.BytesIO(archive_to_excel(ac._load_rows())))["Trades"]
    headers = [c.value for c in sheet[1]]
    pnl_column = headers.index("Realized PnL ($, net of TC)") + 1
    assert sheet.freeze_panes == "A2" and "$" in sheet.cell(row=2, column=pnl_column).number_format
    assert sheet.cell(row=2, column=headers.index("Entry Date") + 1).number_format == "yyyy-mm-dd"

